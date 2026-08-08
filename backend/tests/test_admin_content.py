import io
import os
from uuid import uuid4

from PIL import Image

os.environ.setdefault("DATABASE_URL", "sqlite+pysqlite:///:memory:")
os.environ.setdefault("MEDIA_ROOT", "/tmp/meta-test-media")
os.environ.setdefault("SEED_INITIAL_DATA", "true")

from fastapi.testclient import TestClient
from app.main import app


def admin_headers(client: TestClient) -> dict[str, str]:
    login = client.post("/api/v1/auth/login", json={"email": "admin@example.com", "password": "ChangeMe123!"})
    assert login.status_code == 200, login.text
    return {"Authorization": f"Bearer {login.json()['access_token']}"}


def admin_session(client: TestClient) -> None:
    login = client.post(
        "/admin/login",
        data={"username": "admin@example.com", "password": "ChangeMe123!"},
        follow_redirects=False,
    )
    assert login.status_code in {302, 303}


def jpeg_bytes() -> bytes:
    target = io.BytesIO()
    Image.new("RGB", (32, 32), "white").save(target, format="JPEG")
    return target.getvalue()


def test_admin_dashboard_structure_and_typable_forms():
    with TestClient(app) as client:
        unauthenticated = client.get("/admin/", follow_redirects=True)
        assert unauthenticated.url.path == "/admin/login"
        admin_session(client)
        page = client.get("/admin/management?area=projects&type=tarahi&new=1")
        assert page.status_code == 200
        assert 'name="name"' in page.text
        assert "تصویر کاور پروژه" in page.text
        assert 'name="category_id"' not in page.text
        assert "طراحی" in page.text and "نظارت" in page.text and "اجرا" in page.text
        assert "pointer-events:auto" in page.text
        assert client.get("/docs").status_code == 200


def test_project_admin_create_upload_reference_edit_and_public_api():
    suffix = uuid4().hex[:8]
    with TestClient(app) as client:
        admin_session(client)
        project_slug = f"tarahi-{suffix}"
        created = client.post(
            "/admin/management?area=projects&type=tarahi",
            data={
                "action": "save_project",
                "project_type": "tarahi",
                "name": "Tarahi project",
                "slug": project_slug,
                "metraj": "500",
                "location": "Tehran",
                "description": "API-backed project",
                "state": "open",
                "display_order": "1",
            },
            files={"cover_image": ("cover.jpg", jpeg_bytes(), "image/jpeg")},
            follow_redirects=False,
        )
        assert created.status_code == 303, created.text
        project = next(item for item in client.get("/api/v1/tarahi/projects").json() if item["slug"] == project_slug)

        section = client.post(
            "/admin/management?area=projects&type=tarahi",
            data={
                "action": "save_section",
                "project_type": "tarahi",
                "project_id": project["id"],
                "section_name": "سازه",
                "section_slug": f"structure-{suffix}",
                "section_description": "نقشه‌های سازه پروژه",
                "section_display_order": "1",
            },
            files=[
                ("section_images", ("structure-1.jpg", jpeg_bytes(), "image/jpeg")),
                ("section_images", ("structure-2.jpg", jpeg_bytes(), "image/jpeg")),
            ],
            follow_redirects=False,
        )
        assert section.status_code == 303
        project = next(item for item in client.get("/api/v1/tarahi/projects").json() if item["slug"] == project_slug)
        section_id = project["sections"][0]["id"]

        updated = client.post(
            "/admin/management?area=projects&type=tarahi",
            data={
                "action": "save_project",
                "project_type": "tarahi",
                "project_id": project["id"],
                "name": "Tarahi project",
                "slug": project_slug,
                "metraj": "500",
                "location": "Tehran",
                "description": "API-backed project",
                "state": "open",
                "display_order": "1",
            },
            follow_redirects=False,
        )
        assert updated.status_code == 303, updated.text
        public = client.get("/api/v1/tarahi/projects")
        assert public.status_code == 200
        project = next(item for item in public.json() if item["slug"] == project_slug)
        assert len(project["images"]) == 3
        assert project["sections"][0]["name"] == "سازه"
        assert project["sections"][0]["description"] == "نقشه‌های سازه پروژه"
        assert len(project["sections"][0]["images"]) == 2
        assert all(not image["is_cover"] for image in project["sections"][0]["images"])
        covers = [image for image in project["images"] if image["is_cover"]]
        assert len(covers) == 1
        assert covers[0]["section_id"] is None

        edit = client.get(f"/admin/management?area=projects&type=tarahi&edit={project['id']}")
        assert edit.status_code == 200
        assert "Tarahi project" in edit.text
        assert "انتخاب تصویر کاور" in edit.text
        assert "افزودن تصویر" in edit.text
        assert "حذف تصویر" in edit.text


def test_users_and_outbound_messages_api():
    suffix = uuid4().hex[:8]
    with TestClient(app) as client:
        headers = admin_headers(client)
        created = client.post(
            "/api/v1/admin/users",
            headers=headers,
            json={
                "full_name": "New User",
                "email": f"user-{suffix}@example.com",
                "phone": "09120000000",
                "password": "StrongPass123!",
                "role": "user",
                "is_active": True,
            },
        )
        assert created.status_code == 201, created.text
        user_id = created.json()["id"]
        outgoing = client.post(
            "/api/v1/admin/messages/outgoing",
            headers=headers,
            json={"subject": "Test", "body": "Message body", "channel": "phone", "user_ids": [user_id]},
        )
        assert outgoing.status_code == 201, outgoing.text
        assert outgoing.json()["channel"] == "phone"
        assert outgoing.json()["recipients"][0]["user_id"] == user_id
        listed = client.get("/api/v1/admin/messages/outgoing", headers=headers)
        assert listed.status_code == 200
        assert any(item["id"] == outgoing.json()["id"] for item in listed.json())
        assert client.delete(f"/api/v1/admin/messages/outgoing/{outgoing.json()['id']}", headers=headers).status_code == 204
        assert client.delete(f"/api/v1/admin/users/{user_id}", headers=headers).status_code == 204


def test_admin_users_messages_and_nezarat_table_forms():
    suffix = uuid4().hex[:8]
    with TestClient(app) as client:
        admin_session(client)
        user = client.post(
            "/admin/management?area=users",
            data={
                "action": "save_user",
                "full_name": "Panel User",
                "email": f"panel-{suffix}@example.com",
                "phone": "09121111111",
                "password": "StrongPass123!",
                "role": "user",
                "is_active": "on",
            },
            follow_redirects=False,
        )
        assert user.status_code == 303
        users_page = client.get("/admin/management?area=users")
        assert f"panel-{suffix}@example.com" in users_page.text

        recipient_id = next(item["id"] for item in client.get("/api/v1/admin/users", headers=admin_headers(client)).json() if item["email"] == f"panel-{suffix}@example.com")
        message = client.post(
            "/admin/management?area=messages&message_view=compose",
            data={
                "action": "send_message",
                "recipient_ids": recipient_id,
                "channel": "email",
                "subject": "Panel message",
                "message_body": "Queued from panel",
            },
            follow_redirects=False,
        )
        assert message.status_code == 303
        sent_page = client.get("/admin/management?area=messages&message_view=sent")
        assert "Panel message" in sent_page.text
        assert f"panel-{suffix}@example.com" in sent_page.text

        row = client.post(
            "/admin/management?area=projects&type=nezarat",
            data={
                "action": "save_table_row",
                "project_type": "nezarat",
                "table_status": "live",
                "region": "22",
                "metraj": "800",
                "allowed_floors": "7",
                "project_stage": "اسکلت",
                "table_year": "1405",
                "table_display_order": "1",
                "is_published": "on",
            },
            follow_redirects=False,
        )
        assert row.status_code == 303
        public_rows = client.get("/api/v1/nezarat/table?status=live").json()
        assert any(item["region"] == "22" and item["project_stage"] == "اسکلت" for item in public_rows)


def test_persian_unicode_slugs_are_valid_in_public_responses():
    with TestClient(app) as client:
        admin_session(client)
        response = client.post(
            "/admin/management?area=projects&type=nezarat",
            data={
                "action": "save_relation",
                "project_type": "nezarat",
                "relation_name": "لواسان",
                "relation_slug": "لواسان",
                "relation_display_order": "25",
                "relation_is_active": "on",
            },
            follow_redirects=False,
        )
        assert response.status_code == 303
        public = client.get("/api/v1/nezarat/regions")
        assert public.status_code == 200, public.text
        assert any(item["slug"] == "لواسان" for item in public.json())


def excel_bytes() -> bytes:
    from openpyxl import Workbook

    target = io.BytesIO()
    workbook = Workbook()
    first = workbook.active
    first.title = "پروژه‌های جاری"
    first.append(["منطقه", "متراژ", "مرحله"])
    first.append(["لواسان", 1200, "اسکلت"])
    second = workbook.create_sheet("پایان‌یافته")
    second.append(["نام", "آدرس"])
    second.append(["پروژه الف", "تهران"])
    helper = workbook.create_sheet("Table")
    helper.append(["internal"])
    helper.sheet_state = "hidden"
    workbook.save(target)
    return target.getvalue()


def test_nezarat_excel_upload_is_rendered_exactly_from_api():
    with TestClient(app) as client:
        admin_session(client)
        response = client.post(
            "/admin/management?area=projects&type=nezarat",
            data={"action": "upload_table_workbook", "project_type": "nezarat"},
            files={"excel_file": ("nezarat.xlsx", excel_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            follow_redirects=False,
        )
        assert response.status_code == 303, response.text

        public = client.get("/api/v1/nezarat/table-workbook")
        assert public.status_code == 200, public.text
        payload = public.json()
        assert payload["source_filename"] == "nezarat.xlsx"
        assert [sheet["name"] for sheet in payload["sheets"]] == ["پروژه‌های جاری", "پایان‌یافته"]
        assert payload["sheets"][0]["columns"] == ["منطقه", "متراژ", "مرحله"]
        assert payload["sheets"][0]["rows"][0] == ["لواسان", 1200, "اسکلت"]

        panel = client.get("/admin/management?area=projects&type=nezarat")
        assert 'name="excel_file"' in panel.text
        assert "nezarat.xlsx" in panel.text
        assert "پروژه‌های جاری" in panel.text
