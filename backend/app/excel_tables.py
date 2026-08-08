from __future__ import annotations

import io
import json
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from starlette.datastructures import UploadFile

MAX_EXCEL_BYTES = 10 * 1024 * 1024
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 10_000
MAX_COLUMNS_PER_SHEET = 200
ALLOWED_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xltx", ".xltm"}


def _json_cell(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def parse_excel_upload(upload: UploadFile) -> list[dict[str, Any]]:
    filename = Path(upload.filename or "").name
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_EXCEL_SUFFIXES:
        raise ValueError("فقط فایل‌های Excel با فرمت xlsx یا xlsm قابل بارگذاری هستند.")

    content = upload.file.read(MAX_EXCEL_BYTES + 1)
    if not content:
        raise ValueError("فایل اکسل خالی است.")
    if len(content) > MAX_EXCEL_BYTES:
        raise ValueError("حجم فایل اکسل نباید بیشتر از ۱۰ مگابایت باشد.")

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("فایل اکسل معتبر نیست یا قابل خواندن نیست.") from exc

    sheets: list[dict[str, Any]] = []
    try:
        for worksheet in workbook.worksheets[:MAX_SHEETS]:
            # Helper/hidden worksheets (often named "Table") are not user-facing tabs.
            if worksheet.sheet_state != "visible" or worksheet.title.strip().casefold() == "table":
                continue
            raw_rows: list[list[Any]] = []
            for row_index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if row_index >= MAX_ROWS_PER_SHEET:
                    raise ValueError(f"شیت «{worksheet.title}» بیشتر از {MAX_ROWS_PER_SHEET:,} ردیف دارد.")
                values = [_json_cell(value) for value in row[:MAX_COLUMNS_PER_SHEET]]
                raw_rows.append(values)

            while raw_rows and not any(value not in (None, "") for value in raw_rows[-1]):
                raw_rows.pop()
            if not raw_rows:
                continue

            header_index = next(
                (index for index, row in enumerate(raw_rows) if any(value not in (None, "") for value in row)),
                None,
            )
            if header_index is None:
                continue

            relevant_rows = raw_rows[header_index:]
            last_column = 0
            for row in relevant_rows:
                for index, value in enumerate(row):
                    if value not in (None, ""):
                        last_column = max(last_column, index + 1)
            if last_column == 0:
                continue

            header_values = relevant_rows[0][:last_column]
            columns = [
                str(value).strip() if value not in (None, "") else f"ستون {index + 1}"
                for index, value in enumerate(header_values)
            ]

            rows: list[list[str | int | float | bool | None]] = []
            for raw_row in relevant_rows[1:]:
                normalized = list(raw_row[:last_column])
                if len(normalized) < last_column:
                    normalized.extend([None] * (last_column - len(normalized)))
                # Preserve internal blank rows, but trailing blank rows were already removed.
                rows.append(normalized)

            sheets.append({"name": worksheet.title, "columns": columns, "rows": rows})
    finally:
        workbook.close()

    if not sheets:
        raise ValueError("هیچ جدول قابل نمایش در فایل اکسل پیدا نشد.")
    return sheets


def dumps_sheets(sheets: list[dict[str, Any]]) -> str:
    return json.dumps(sheets, ensure_ascii=False, separators=(",", ":"))


def loads_sheets(raw: str | None) -> list[dict[str, Any]]:
    try:
        value = json.loads(raw or "[]")
    except json.JSONDecodeError:
        return []
    if not isinstance(value, list):
        return []
    return [
        sheet
        for sheet in value
        if isinstance(sheet, dict)
        and str(sheet.get("name", "")).strip().casefold() != "table"
    ]
